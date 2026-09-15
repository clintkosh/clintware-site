from __future__ import annotations

from html import escape
from pathlib import Path
from typing import TYPE_CHECKING

from PySide6.QtCore import Qt

if TYPE_CHECKING:
    from PySide6.QtWidgets import QTextEdit


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def load_docx(path: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(path)
    chunks: list[str] = []
    align_map = {
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.LEFT: "left",
    }
    for p in doc.paragraphs:
        style_name = (p.style.name if p.style else "") or ""
        tag = "p"
        if style_name.lower().startswith("heading"):
            level = "".join(ch for ch in style_name if ch.isdigit()) or "2"
            tag = f"h{min(max(int(level), 1), 6)}"
        align = align_map.get(p.alignment)
        attrs = f' style="text-align:{align}"' if align else ""
        runs = []
        for run in p.runs:
            text = escape(run.text).replace("\n", "<br>")
            if not text:
                continue
            if run.bold:
                text = f"<strong>{text}</strong>"
            if run.italic:
                text = f"<em>{text}</em>"
            if run.underline:
                text = f"<u>{text}</u>"
            runs.append(text)
        chunks.append(f"<{tag}{attrs}>{''.join(runs) or '<br>'}</{tag}>")
    return "\n".join(chunks)


def save_docx(editor: "QTextEdit", path: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    qt_doc = editor.document()
    block = qt_doc.begin()
    align_map = {
        Qt.AlignmentFlag.AlignCenter: WD_ALIGN_PARAGRAPH.CENTER,
        Qt.AlignmentFlag.AlignRight: WD_ALIGN_PARAGRAPH.RIGHT,
        Qt.AlignmentFlag.AlignJustify: WD_ALIGN_PARAGRAPH.JUSTIFY,
        Qt.AlignmentFlag.AlignLeft: WD_ALIGN_PARAGRAPH.LEFT,
    }

    while block.isValid():
        fmt = block.blockFormat()
        para = document.add_paragraph()
        alignment = fmt.alignment()
        for qt_flag, wd_flag in align_map.items():
            if alignment & qt_flag:
                para.alignment = wd_flag
                break

        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid() and frag.text():
                run = para.add_run(frag.text())
                cf = frag.charFormat()
                run.bold = cf.fontWeight() >= 700
                run.italic = cf.fontItalic()
                run.underline = cf.fontUnderline()
                if cf.fontPointSize() > 0:
                    run.font.size = Pt(cf.fontPointSize())
                family = cf.fontFamily()
                if family:
                    run.font.name = family
            it += 1
        block = block.next()

    document.save(path)


def load_text_like(path: str) -> tuple[str, bool]:
    suffix = Path(path).suffix.lower()
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    return text, suffix in {".html", ".htm"}


def save_text_like(editor: "QTextEdit", path: str) -> None:
    suffix = Path(path).suffix.lower()
    if suffix in {".html", ".htm"}:
        Path(path).write_text(editor.document().toHtml(), encoding="utf-8")
    else:
        Path(path).write_text(editor.toPlainText(), encoding="utf-8")
